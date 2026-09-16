"""
Importação do DRE por centro de custo exportado do Contimatic.

O relatório vem como um único bloco em cascata, uma seção por obra:

    OBRA: 318 VISTA BROOKLIN - ADOLPHO
      Receitas Brutas
        Servicos prestados - mercado interno - 25      217.704,60
      Receitas Brutas Total...                         217.704,60
      Deduções
        INSS S/ FATURAMENTO - 544                       (5.878,02)
      Deduções Total...                                 (5.878,02)
      = Receita Líquida                                211.826,58
      Custos
        Salarios e Ordenados - 311                     (28.270,19)
        ...
      Custos Total...                                  (95.526,97)
      = Lucro Bruto                                    116.299,61

O módulo está partido em duas camadas de propósito:

  * **ler_relatorio(caminho)** — conhece o layout do arquivo e devolve linhas
    normalizadas. É a única parte que depende do formato.

  * **importar_linhas(...)** — resolve obra e conta, atualiza o DRE e guarda o
    detalhe. Não sabe nada sobre Excel, e por isso é testável sem arquivo.

**O relatório não traz a competência.** Não há no arquivo nada que diga a que
mês os valores pertencem, então quem importa informa o mês e o ano na tela de
envio, e é isso que decide onde os valores são gravados.

Uma linha normalizada é um dicionário:

    {
        "obra_codigo":  "318",
        "obra_nome":    "VISTA BROOKLIN - ADOLPHO",
        "secao":        "receita",        # ver SECOES
        "conta_codigo": "25",
        "conta_nome":   "Servicos prestados - mercado interno",
        "valor":        217704.60,        # sinal como veio do relatório
        "linha":        12,               # linha da planilha, para diagnóstico
    }
"""

import re
import difflib
import datetime
import unicodedata

import openpyxl


# Seções do relatório e como cada uma entra no DRE do sistema.
#
# 'receita' e 'custo' têm equivalente direto no plano de contas. As outras duas
# não: o DRE do sistema calcula imposto e despesa administrativa por alíquota
# (13,15% e 5,56%), enquanto o Contimatic traz o valor realmente contabilizado.
# Importar as duas coisas somaria a mesma despesa duas vezes, então essas contas
# ficam pendentes de decisão em vez de entrar sozinhas.
SECOES = {
    "receitas brutas": "receita",
    "receita bruta": "receita",
    "outras receitas": "receita",
    "deducoes": "deducao",
    "custos": "custo",
    "custo": "custo",
    "despesas administrativas": "despesa_administrativa",
    "despesas operacionais": "despesa_administrativa",
    "despesas financeiras": "despesa_financeira",
}

SECOES_QUE_VIRAM_LANCAMENTO = {"receita", "custo"}

TIPO_DA_SECAO = {"receita": "receita", "custo": "custo"}

# Em que tipo do plano de contas a sugestão pode procurar, por seção.
# Dedução e despesa são redutoras: a sugestão nunca pode atravessar para
# receita. Sem essa trava, "Serviços prestados por terceiros" (uma despesa)
# era sugerida como "Serviços prestados" (uma receita), o que transformaria
# gasto em faturamento.
TIPO_SUGERIDO_DA_SECAO = {
    "receita": "receita",
    "custo": "custo",
    "deducao": "custo",
    "despesa_administrativa": "custo",
    "despesa_financeira": "custo",
}

ROTULO_DA_SECAO = {
    "receita": "Receitas Brutas",
    "deducao": "Deduções",
    "custo": "Custos",
    "despesa_administrativa": "Despesas Administrativas",
    "despesa_financeira": "Despesas Financeiras",
}


def _sem_acento(texto):
    nfkd = unicodedata.normalize("NFKD", str(texto if texto is not None else ""))
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def _normalizar(texto):
    return " ".join(_sem_acento(texto).lower().split())


def _codigo_limpo(codigo):
    """'0000311' e '311' são a mesma conta; '3.1.1' também aparece."""
    if codigo is None:
        return ""
    texto = str(codigo).strip()
    if texto.endswith(".0"):          # o Excel entrega número como float
        texto = texto[:-2]
    somente_digitos = "".join(c for c in texto if c.isdigit())
    return somente_digitos.lstrip("0") or somente_digitos


def chave_da_conta(conta_codigo, conta_nome):
    """
    Identificador da conta no de-para.

    Nem toda linha do relatório traz código — a coluna estreita corta o nome e
    às vezes o código junto ("Locação de Maqs, Ferramentas e Equipamen"). Sem
    uma chave, essas contas não conseguiam nem virar pendência e ficariam fora
    do DRE para sempre, sem ninguém ver. Quando não há código, o próprio nome
    normalizado vira a chave.
    """
    codigo = _codigo_limpo(conta_codigo)
    if codigo:
        return codigo
    nome = _normalizar(conta_nome)
    return f"nome:{nome}" if nome else ""


# ---------------------------------------------------------------------------
# Leitura do arquivo
# ---------------------------------------------------------------------------

# "OBRA: 251 ESTHER TOERS", "OBRA: 319 - BOSQUE VILA NOVA - R YAZBEK"
PADRAO_OBRA = re.compile(r"^obra\s*:?\s*(\d+)\s*-?\s*(.*)$", re.IGNORECASE)

# O nome da conta pode conter hífen ("Servicos prestados - mercado interno - 25"),
# então o código é o ÚLTIMO " - <números>" da linha.
PADRAO_CONTA = re.compile(r"^(.*?)\s*-\s*(\d+)\s*$")

PADRAO_TOTAL = re.compile(r"total\s*\.*\s*$", re.IGNORECASE)


def _valor_da_celula(valor):
    """
    Devolve float ou None.

    O Excel normalmente entrega número, mas o relatório pode vir com a coluna
    como texto — inclusive no formato contábil, em que o negativo aparece entre
    parênteses: "(7.936,28)".
    """
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)

    texto = str(valor).strip()
    if not texto:
        return None

    negativo = texto.startswith("(") and texto.endswith(")")
    texto = texto.strip("()").replace("R$", "").strip()
    texto = re.sub(r"[^\d,.\-]", "", texto)
    if not texto or texto in {"-", ".", ","}:
        return None

    # Com os dois separadores, o último é o decimal.
    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(",", ".")

    try:
        numero = float(texto)
    except ValueError:
        return None

    return -numero if negativo else numero


def _partes_da_linha(celulas):
    """
    Separa a linha em (rótulo, valor).

    Não assume posição fixa de coluna: o rótulo é o primeiro texto e o valor é o
    último número da linha. Isso aguenta o relatório vir com colunas a mais, ou
    com a primeira coluna usada só para recuo.
    """
    rotulo = None
    valor = None

    for celula in celulas:
        if celula is None:
            continue
        if rotulo is None and isinstance(celula, str) and celula.strip():
            rotulo = celula.strip()
            continue
        numero = _valor_da_celula(celula)
        if numero is not None:
            valor = numero

    return rotulo, valor


def _identificar_obra(rotulo):
    """Devolve (codigo, nome) se o rótulo for um cabeçalho de centro de custo."""
    normalizado = _normalizar(rotulo)

    achado = PADRAO_OBRA.match(rotulo.strip())
    if achado:
        nome = achado.group(2).strip(" -")
        return achado.group(1), nome or f"OBRA {achado.group(1)}"

    if "demais obras" in normalizado:
        return "DEMAIS", "DEMAIS OBRAS"

    if re.fullmatch(r"(depto|departamento)\s+tecnico", normalizado):
        return "DEPTO-TEC", "Departamento Técnico (Administrativo)"

    return None, None


def ler_relatorio(caminho_arquivo, aba=None):
    """
    Lê o relatório e devolve:

        {
          "linhas":  [linha normalizada, ...],
          "obras":   {codigo: nome},
          "totais":  {(codigo, secao): valor declarado pelo próprio relatório},
          "avisos":  [texto, ...],
        }

    Os totais declarados ("Custos Total...") são guardados para conferência:
    depois de somar as contas lidas, o sistema compara com o que o relatório
    afirma. Divergiu, a leitura errou em algum lugar — e é melhor saber disso
    antes de gravar.
    """
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    ws = wb[aba] if aba else wb[wb.sheetnames[0]]

    linhas = []
    obras = {}
    totais = {}
    avisos = []

    obra_codigo = None
    obra_nome = None
    secao = None

    for numero, celulas in enumerate(ws.iter_rows(values_only=True), start=1):
        rotulo, valor = _partes_da_linha(celulas)
        if not rotulo:
            continue

        texto = rotulo.strip()
        normalizado = _normalizar(texto)

        # 1. Cabeçalho de centro de custo
        codigo, nome = _identificar_obra(texto)
        if codigo:
            obra_codigo, obra_nome = codigo, nome
            obras.setdefault(codigo, nome)
            secao = None
            continue

        # 2. Linha de resultado ("= Lucro Bruto", "= Prejuízo") — só fecha a seção
        if texto.startswith("="):
            secao = None
            continue

        # 3. Linha de total da seção — guarda para conferir depois
        if PADRAO_TOTAL.search(normalizado):
            rotulo_secao = PADRAO_TOTAL.sub("", normalizado).strip(" .")
            chave_secao = SECOES.get(rotulo_secao)
            if obra_codigo and chave_secao and valor is not None:
                totais[(obra_codigo, chave_secao)] = valor
            secao = None
            continue

        # 4. Cabeçalho de seção
        if normalizado in SECOES:
            secao = SECOES[normalizado]
            continue

        # 5. Linha de conta
        if secao is None or obra_codigo is None or valor is None:
            continue

        achado = PADRAO_CONTA.match(texto)
        if achado:
            conta_nome, conta_codigo = achado.group(1).strip(), achado.group(2)
        else:
            # Acontece quando o nome vem cortado pela largura da coluna, às
            # vezes com o hífen solto no fim: "Despesas com Radio e Telefonia -"
            conta_nome, conta_codigo = texto.rstrip(" -"), None

        linhas.append({
            "obra_codigo": obra_codigo,
            "obra_nome": obra_nome,
            "secao": secao,
            "conta_codigo": conta_codigo,
            "conta_nome": conta_nome,
            "valor": valor,
            "linha": numero,
        })

    if not linhas:
        avisos.append(
            "Nenhuma conta foi reconhecida. Confira se o arquivo é o DRE por centro "
            "de custo do Contimatic e se a aba certa foi escolhida."
        )

    return {"linhas": linhas, "obras": obras, "totais": totais, "avisos": avisos}


def conferir_totais(leitura, tolerancia=0.01):
    """
    Compara a soma das contas lidas com os totais que o próprio relatório
    declara. Divergência aponta linha não lida ou lida duas vezes.
    """
    somado = {}
    for linha in leitura["linhas"]:
        chave = (linha["obra_codigo"], linha["secao"])
        somado[chave] = somado.get(chave, 0.0) + linha["valor"]

    divergencias = []
    for chave, declarado in leitura["totais"].items():
        calculado = somado.get(chave, 0.0)
        if abs(calculado - declarado) > tolerancia:
            divergencias.append({
                "obra_codigo": chave[0],
                "secao": chave[1],
                "declarado": declarado,
                "lido": calculado,
                "diferenca": calculado - declarado,
            })

    return sorted(divergencias, key=lambda d: -abs(d["diferenca"]))


# ---------------------------------------------------------------------------
# Resolução de conta contábil -> categoria do plano de contas
# ---------------------------------------------------------------------------

class ResolvedorDeContas:
    """
    Descobre a que categoria do plano de contas pertence cada conta contábil.

    Ordem: de-para explícito (contas_map), código do plano, nome exato. O que
    não casar volta como pendência, com uma sugestão por semelhança de nome
    para a pessoa confirmar.

    A sugestão nunca é aplicada sozinha: uma conta no lugar errado desloca
    dinheiro entre linhas do DRE sem dar erro nenhum.
    """

    SEMELHANCA_MINIMA = 0.78

    def __init__(self, cur):
        self.cur = cur

        cur.execute("SELECT id, codigo, nome, tipo FROM categorias_conta")
        categorias = cur.fetchall()
        self.categorias = [dict(r) for r in categorias]
        self.por_codigo = {_codigo_limpo(r["codigo"]): r["id"] for r in categorias if r["codigo"]}
        self.por_nome = {_normalizar(r["nome"]): r["id"] for r in categorias}
        self.nome_por_id = {r["id"]: r["nome"] for r in categorias}

        cur.execute("SELECT conta_codigo, categoria_id, ignorar FROM contas_map")
        # A chave vem gravada pronta em contas_map: código quando existe,
        # "nome:<normalizado>" quando a conta veio sem código no relatório.
        self.mapa = {
            str(r["conta_codigo"]): (r["categoria_id"], r["ignorar"])
            for r in cur.fetchall()
        }

        self.pendencias = {}

    def resolver(self, conta_codigo, conta_nome):
        """Devolve (categoria_id, motivo). categoria_id None = não resolvida."""
        codigo = _codigo_limpo(conta_codigo)
        chave = chave_da_conta(conta_codigo, conta_nome)

        if chave and chave in self.mapa:
            categoria_id, ignorar = self.mapa[chave]
            if ignorar:
                return None, "ignorada"
            if categoria_id:
                return categoria_id, "de-para"

        if codigo and codigo in self.por_codigo:
            return self.por_codigo[codigo], "codigo"

        nome = _normalizar(conta_nome)
        if nome and nome in self.por_nome:
            return self.por_nome[nome], "nome"

        return None, "pendente"

    def sugerir(self, conta_nome, tipo_esperado=None):
        """Categoria mais parecida pelo nome, para a tela de revisão propor."""
        alvo = _normalizar(conta_nome)
        if not alvo:
            return None

        candidatas = [
            c for c in self.categorias
            if tipo_esperado is None or c["tipo"] == tipo_esperado
        ]
        nomes = {_normalizar(c["nome"]): c for c in candidatas}

        # O Contimatic costuma detalhar o que o plano de contas resume
        # ("Servicos prestados - mercado interno" contra "Serviços prestados"),
        # e a coluna estreita corta o nome no fim ("...e Equipamen"). Nos dois
        # casos um nome é começo do outro, o que a distância de texto pura não
        # captura bem — então o prefixo é testado antes.
        melhor_prefixo = None
        for nome_categoria in nomes:
            if not nome_categoria:
                continue
            curto, longo = sorted((nome_categoria, alvo), key=len)
            if len(curto) >= 8 and longo.startswith(curto):
                if melhor_prefixo is None or len(curto) > len(melhor_prefixo):
                    melhor_prefixo = nome_categoria

        if melhor_prefixo:
            categoria = nomes[melhor_prefixo]
            return {
                "categoria_id": categoria["id"],
                "categoria_nome": categoria["nome"],
                "semelhanca": difflib.SequenceMatcher(None, alvo, melhor_prefixo).ratio(),
                "por_prefixo": True,
            }

        proximos = difflib.get_close_matches(
            alvo, nomes.keys(), n=1, cutoff=self.SEMELHANCA_MINIMA
        )
        if not proximos:
            return None

        categoria = nomes[proximos[0]]
        return {
            "categoria_id": categoria["id"],
            "categoria_nome": categoria["nome"],
            "semelhanca": difflib.SequenceMatcher(None, alvo, proximos[0]).ratio(),
            "por_prefixo": False,
        }

    def registrar_pendencia(self, linha, motivo):
        chave = chave_da_conta(linha["conta_codigo"], linha["conta_nome"])
        pendencia = self.pendencias.setdefault(chave, {
            "chave": chave,
            "conta_codigo": linha["conta_codigo"],
            "conta_nome": linha["conta_nome"],
            "secao": linha["secao"],
            "motivo": motivo,
            "ocorrencias": 0,
            "valor_total": 0.0,
            "obras": set(),
            "sugestao": self.sugerir(
                linha["conta_nome"], TIPO_SUGERIDO_DA_SECAO.get(linha["secao"])
            ),
        })
        pendencia["ocorrencias"] += 1
        pendencia["valor_total"] += abs(linha["valor"] or 0)
        pendencia["obras"].add(linha["obra_codigo"])
        return pendencia


def resolver_obras(cur):
    """Código do centro de custo -> id da obra."""
    cur.execute("SELECT id, codigo FROM obras")
    mapa = {}
    for r in cur.fetchall():
        mapa[_codigo_limpo(r["codigo"])] = r["id"]
        mapa[str(r["codigo"]).strip().upper()] = r["id"]
    return mapa


# ---------------------------------------------------------------------------
# Gravação
# ---------------------------------------------------------------------------

def importar_linhas(cur, linhas, competencia, registro, importacao_id=None,
                    criar_obras=True, empresa_id=None):
    """
    Grava as linhas do relatório na competência informada.

    'competencia' é (ano, mes) e vem da tela, porque o relatório não traz data.

    Só as seções de receita e custo viram lançamento. Dedução e despesa
    administrativa ficam pendentes: o DRE do sistema já calcula essas duas por
    alíquota, e gravar também o valor contabilizado cobraria a mesma despesa
    duas vezes. Quem importa decide, uma vez, na tela de revisão.
    """
    ano, mes = competencia
    resolvedor = ResolvedorDeContas(cur)
    obras_por_codigo = resolver_obras(cur)

    agora = datetime.datetime.now().isoformat()

    celulas = {}
    detalhes = []
    obras_criadas = []
    obras_desconhecidas = {}
    ignoradas = 0

    for linha in linhas:
        codigo_obra = _codigo_limpo(linha["obra_codigo"]) or str(linha["obra_codigo"]).upper()
        obra_id = obras_por_codigo.get(codigo_obra)

        if not obra_id:
            if criar_obras and empresa_id:
                cur.execute(
                    """
                    INSERT INTO obras (empresa_id, nome, codigo, status)
                    VALUES (?, ?, ?, 'em_andamento')
                    """,
                    (empresa_id,
                     linha["obra_nome"] or f"OBRA {linha['obra_codigo']}",
                     str(linha["obra_codigo"])),
                )
                obra_id = cur.lastrowid
                obras_por_codigo[codigo_obra] = obra_id
                obras_criadas.append(f"{linha['obra_codigo']} - {linha['obra_nome']}")
            else:
                pendente = obras_desconhecidas.setdefault(
                    str(linha["obra_codigo"]), {"nome": linha["obra_nome"], "ocorrencias": 0}
                )
                pendente["ocorrencias"] += 1
                continue

        # A decisão de quem importa vem primeiro. Mandar ignorar ou apontar uma
        # categoria vale para qualquer seção — é essa a saída para dedução e
        # despesa administrativa, que sozinhas não entram.
        categoria_id, motivo = resolvedor.resolver(linha["conta_codigo"], linha["conta_nome"])

        if motivo == "ignorada":
            ignoradas += 1
            continue

        if not categoria_id:
            if linha["secao"] not in SECOES_QUE_VIRAM_LANCAMENTO:
                resolvedor.registrar_pendencia(linha, "secao_sem_equivalente")
            else:
                resolvedor.registrar_pendencia(linha, motivo)
            ignoradas += 1
            continue

        if motivo != "de-para" and linha["secao"] not in SECOES_QUE_VIRAM_LANCAMENTO:
            # Casou por código ou nome, mas a seção não tem equivalente no DRE.
            # Precisa de confirmação explícita, senão a despesa entra duas vezes:
            # uma pelo valor contabilizado e outra pela alíquota que o DRE aplica.
            resolvedor.registrar_pendencia(linha, "secao_sem_equivalente")
            ignoradas += 1
            continue

        chave = (obra_id, categoria_id)
        celulas[chave] = celulas.get(chave, 0.0) + abs(linha["valor"] or 0)

        detalhes.append({
            "obra_id": obra_id,
            "categoria_id": categoria_id,
            "conta_codigo": linha["conta_codigo"],
            "conta_nome": linha["conta_nome"],
            "secao": linha["secao"],
            "valor": abs(linha["valor"] or 0),
        })

    # O detalhe da competência sai antes de entrar o novo: o relatório do mês é
    # a verdade, e reimportar não pode duplicar nem deixar resto do envio anterior.
    for obra_id in {d["obra_id"] for d in detalhes}:
        cur.execute(
            "DELETE FROM partidas WHERE obra_id = ? AND ano = ? AND mes = ?",
            (obra_id, ano, mes),
        )

    for d in detalhes:
        cur.execute(
            """
            INSERT INTO partidas
                (obra_id, categoria_id, mes, ano, data, documento, historico,
                 conta_codigo, conta_nome, secao, valor, importacao_id)
            VALUES (?, ?, ?, ?, NULL, NULL, NULL, ?, ?, ?, ?, ?)
            """,
            (d["obra_id"], d["categoria_id"], mes, ano,
             d["conta_codigo"], d["conta_nome"], d["secao"], d["valor"], importacao_id),
        )

    # O total passa pelo RegistroImportacao, então continua valendo a proteção
    # do lançamento manual e o desfazer.
    gravados = 0
    for (obra_id, categoria_id), total in celulas.items():
        if registro.gravar_lancamento(obra_id, categoria_id, mes, ano, total, agora):
            gravados += 1

    pendencias = []
    for p in resolvedor.pendencias.values():
        p = dict(p)
        p["obras"] = sorted(p["obras"])
        pendencias.append(p)
    pendencias.sort(key=lambda p: -p["valor_total"])

    return {
        "competencia": (ano, mes),
        "linhas_lidas": len(linhas),
        "celulas_atualizadas": gravados,
        "partidas_gravadas": len(detalhes),
        "linhas_ignoradas": ignoradas,
        "obras_criadas": obras_criadas,
        "obras_desconhecidas": obras_desconhecidas,
        "pendencias": pendencias,
    }


def buscar_partidas(cur, obra_id, categoria_id, ano, mes):
    """O detalhe por trás de uma célula do DRE, para a tela de conferência."""
    cur.execute(
        """
        SELECT data, documento, historico, conta_codigo, conta_nome, secao, valor
        FROM partidas
        WHERE obra_id = ? AND categoria_id = ? AND ano = ? AND mes = ?
        ORDER BY id
        """,
        (obra_id, categoria_id, ano, mes),
    )
    return [dict(r) for r in cur.fetchall()]
